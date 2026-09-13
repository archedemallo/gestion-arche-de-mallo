#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Migration des données Suivi (export Excel des Google Sheets) vers le schéma
Supabase créé par 001_schema_arche_de_mallo.sql.

USAGE :
  1. Adapter XLSX_PATH et DATABASE_URL ci-dessous.
     DATABASE_URL se trouve dans Supabase > Project Settings > Database >
     Connection string (utiliser la connexion directe, pas le pooler, pour
     une migration ponctuelle).
  2. Exécuter le schéma 001_schema_arche_de_mallo.sql sur la base cible AVANT
     de lancer ce script (les tables doivent déjà exister).
  3. python3 migrate_suivi.py

CE QUE CE SCRIPT FAIT :
  - Lit chaque onglet peuplé de l'export Suivi.
  - Déduplique les personnes (par nom+prénom) et les animaux (par nom_usuel)
    au fil de l'eau : une même personne/animal mentionné dans plusieurs
    onglets n'est inséré qu'une seule fois.
  - Crée des comptes utilisateurs "placeholder" à partir du champ saisiPar
    (ex: JONATHAN, MALLO...) avec un email temporaire @arche-import.local —
    à FAIRE POINTER vers les vrais comptes une fois l'authentification en
    place (voir avertissement affiché en fin d'exécution).

CE QUE CE SCRIPT NE FAIT PAS (hors périmètre de cette première passe) :
  - Les onglets vides dans l'export ne sont pas traités (rien à migrer).
  - Les fichiers (photos, PDF, certificats ICAD en base64 dans l'onglet
    Divers) ne sont pas migrés : c'est une migration de stockage de fichiers
    (vers Supabase Storage), pas une migration de données de table.
  - Les dates de vermifuge/antipuce/stérilisation *prévues* de l'onglet
    Adoption ne sont pas importées dans animaux_soins_veto : ce sont des
    échéances planifiées, pas des soins effectués, donc pas la même donnée.
  - Le Compta n'est pas couvert : aucun export réel n'a été fourni pour ce
    projet dans cette conversation.
"""

import re
import sys
from datetime import datetime, date

import openpyxl
import psycopg2

XLSX_PATH = "Suivi_des_formulaires_Arche.xlsx"
DATABASE_URL = "postgresql://postgres:PASSWORD@HOST:5432/postgres"  # <-- à remplacer


# ============================================================
# Nettoyage / normalisation
# ============================================================

def clean_str(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def clean_date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def clean_num_text(v):
    """Pour les champs numériques-en-texte (téléphone, code postal, n° chèque,
    puce...) : Excel/Sheets les a convertis en nombre à un moment, avec perte
    du zéro de tête déjà actée dans la donnée SOURCE (pas récupérable ici,
    c'est une limite de l'export, pas de ce script)."""
    if v is None:
        return None
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return str(v)
    return clean_str(v)


def clean_amount(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".").replace("€", "").strip()
    try:
        return float(s) if s else None
    except ValueError:
        return None


def norm_key(*parts):
    return "|".join((p or "").strip().upper() for p in parts)


def slugify_email(name):
    s = re.sub(r"[^a-zA-Z0-9]+", ".", (name or "inconnu").strip().lower()).strip(".")
    return f"{s or 'inconnu'}@arche-import.local"


# ============================================================
# Lecture générique d'un onglet (liste de dicts, lignes vides ignorées)
# ============================================================

def read_sheet(wb, name):
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        d = {}
        for h, v in zip(headers, row):
            # en cas d'en-tête dupliqué dans l'onglet source, la dernière
            # colonne du même nom écrase silencieusement les précédentes —
            # sans impact ici car les onglets concernés sont vides dans cet
            # export (voir README du script).
            d[h] = v
        out.append(d)
    return out


# ============================================================
# Registre get-or-create (personnes / animaux / utilisateurs)
# ============================================================

class Registry:
    def __init__(self, cur):
        self.cur = cur
        self.personnes = {}
        self.animaux = {}
        self.utilisateurs = {}
        self.stats = {
            "personnes_creees": 0, "personnes_reutilisees": 0,
            "animaux_crees": 0, "animaux_reutilises": 0,
            "utilisateurs_crees": 0,
        }

    def get_or_create_utilisateur(self, nom):
        nom = clean_str(nom)
        if not nom:
            return None
        key = norm_key(nom)
        if key in self.utilisateurs:
            return self.utilisateurs[key]
        self.cur.execute(
            "insert into utilisateurs (nom, email) values (%s, %s) "
            "on conflict (email) do update set nom = excluded.nom "
            "returning id",
            (nom, slugify_email(nom)),
        )
        uid = self.cur.fetchone()[0]
        self.utilisateurs[key] = uid
        self.stats["utilisateurs_crees"] += 1
        return uid

    def get_or_create_personne(self, nom=None, prenom=None, raison_sociale=None,
                                representant=None, adresse=None, code_postal=None,
                                ville=None, email=None, telephone=None,
                                date_naissance=None):
        nom, prenom, raison_sociale = clean_str(nom), clean_str(prenom), clean_str(raison_sociale)
        if not (nom or raison_sociale):
            return None
        key = norm_key(nom or raison_sociale, prenom)
        if key in self.personnes:
            self.stats["personnes_reutilisees"] += 1
            return self.personnes[key]
        self.cur.execute(
            """
            insert into personnes (type_personne, nom, prenom, raison_sociale,
                representant, adresse, code_postal, ville, email, telephone,
                date_naissance)
            values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            returning id
            """,
            ("entreprise" if raison_sociale else "particulier",
             nom, prenom, raison_sociale, clean_str(representant),
             clean_str(adresse), clean_num_text(code_postal), clean_str(ville),
             clean_str(email), clean_num_text(telephone), clean_date(date_naissance)),
        )
        pid = self.cur.fetchone()[0]
        self.personnes[key] = pid
        self.stats["personnes_creees"] += 1
        return pid

    def get_or_create_animal(self, nom_usuel, **kw):
        nom_usuel = clean_str(nom_usuel)
        if not nom_usuel:
            return None
        key = norm_key(nom_usuel)
        if key in self.animaux:
            self.stats["animaux_reutilises"] += 1
            return self.animaux[key]
        self.cur.execute(
            """
            insert into animaux (nom_usuel, nom_adoption, sexe, couleur,
                date_naissance, puce, nom_maman, signes_particuliers, espece)
            values (%s,%s,%s,%s,%s,%s,%s,%s,'chat')
            returning id
            """,
            (nom_usuel, clean_str(kw.get("nom_adoption")), clean_str(kw.get("sexe")),
             clean_str(kw.get("couleur")), clean_date(kw.get("date_naissance")),
             clean_num_text(kw.get("puce")), clean_str(kw.get("nom_maman")),
             clean_str(kw.get("signes_particuliers"))),
        )
        aid = self.cur.fetchone()[0]
        self.animaux[key] = aid
        self.stats["animaux_crees"] += 1
        return aid


# ============================================================
# Chèques multiples (pattern commun Dons / Adoptions / Réservations)
# ============================================================

def insert_cheques(cur, table, fk_col, parent_id, row):
    """cheque1..4 / montant_cheque1..4, + le couple simple numeroPaiement/
    numeroCheque + montant_cheque s'il n'y a qu'un seul chèque."""
    any_multi = False
    for i in range(1, 5):
        num = row.get(f"cheque{i}")
        montant = row.get(f"montant_cheque{i}")
        if num is not None or montant is not None:
            any_multi = True
            cur.execute(
                f"insert into {table} ({fk_col}, numero_cheque, montant) values (%s,%s,%s)",
                (parent_id, clean_num_text(num), clean_amount(montant)),
            )
    if not any_multi:
        num = row.get("numeroPaiement") or row.get("numeroCheque")
        montant = row.get("montant_cheque")
        if num is not None or montant is not None:
            cur.execute(
                f"insert into {table} ({fk_col}, numero_cheque, montant) values (%s,%s,%s)",
                (parent_id, clean_num_text(num), clean_amount(montant)),
            )


# ============================================================
# Migration par onglet
# ============================================================

def migrate_adhesion(cur, reg, wb):
    n = 0
    for r in read_sheet(wb, "Adhesion"):
        pid = reg.get_or_create_personne(
            nom=r.get("nom"), prenom=r.get("prenom"), adresse=r.get("adresse"),
            code_postal=r.get("codePostal"), ville=r.get("ville"),
            email=r.get("email"), telephone=r.get("portable"),
            date_naissance=r.get("dateNaissanceMembre"),
        )
        uid = reg.get_or_create_utilisateur(r.get("saisiPar"))
        cur.execute(
            """
            insert into adhesions (personne_id, date_adhesion, type_adhesion,
                cotisation_adherent, don_bienfaiteur, don_sympathisant,
                montant_total, saisi_par, date_fait)
            values (%s,%s,%s,%s,%s,%s,%s,%s,%s)
            """,
            (pid, clean_date(r.get("dateAdhesion")), clean_str(r.get("typesAdhesion")),
             clean_amount(r.get("cotisationAdherent")), clean_amount(r.get("donBienfaiteur")),
             clean_amount(r.get("donSympatisant")), clean_amount(r.get("montantAdhesion")),
             uid, clean_date(r.get("dateFait"))),
        )
        # NB : la feuille Adhesion (COLONNES_FIXES) ne stockait qu'un résumé
        # texte du mode de règlement (check_paiement), pas le montant précis
        # par mode (virement/espèce) — cette précision n'existe donc pas
        # dans l'historique et ne peut pas être reconstituée ici. Les
        # colonnes montant_virement/montant_espece restent NULL pour les
        # lignes migrées ; seuls les nouveaux formulaires les rempliront.
        n += 1
    return n


def migrate_dons(cur, reg, wb):
    n = 0
    for r in read_sheet(wb, "Dons"):
        pid = reg.get_or_create_personne(
            nom=r.get("nom"), prenom=r.get("prenom"), raison_sociale=r.get("raisonSociale"),
            representant=r.get("representant"), adresse=r.get("adresse"),
            code_postal=r.get("codePostal"), ville=r.get("ville"),
            email=r.get("email"), telephone=r.get("telephone"),
        )
        uid = reg.get_or_create_utilisateur(r.get("saisiPar"))
        cur.execute(
            """
            insert into dons (personne_id, don_anonyme, date_signature, montant,
                montant_virement, montant_espece, montant_cb, numero_recu_fiscal, saisi_par)
            values (%s,%s,%s,%s,%s,%s,%s,%s,%s)
            returning id
            """,
            (pid, pid is None, clean_date(r.get("dateSignature")),
             clean_amount(r.get("montantDon")), clean_amount(r.get("montant_virement")),
             clean_amount(r.get("montant_espece")), clean_amount(r.get("montant_cb")),
             clean_num_text(r.get("numeroRecuFiscal")), uid),
        )
        don_id = cur.fetchone()[0]
        # Un chèque unique (numeroCheque/montant_cheque) est traité comme une
        # ligne dons_cheques à 1 élément par insert_cheques() (repli déjà
        # prévu dans cette fonction), donc numero_cheque n'a plus besoin
        # d'exister comme colonne scalaire séparée sur dons.
        insert_cheques(cur, "dons_cheques", "don_id", don_id, r)
        n += 1
    return n


def migrate_depot_chat(cur, reg, wb):
    n = 0
    for r in read_sheet(wb, "Depot_Chat"):
        pid = reg.get_or_create_personne(
            nom=r.get("nom"), prenom=r.get("prenom"), adresse=r.get("adresse"),
            code_postal=r.get("codePostal"), ville=r.get("ville"),
            email=r.get("email"), telephone=r.get("telephone"),
        )
        uid = reg.get_or_create_utilisateur(r.get("saisiPar"))
        cur.execute(
            """
            insert into depots_chat (personne_id, date_depot, commune, preciser_autre,
                nombre_chats_historique, nombre_chatons_historique,
                nombre_autres_historique, saisi_par)
            values (%s,%s,%s,%s,%s,%s,%s,%s)
            """,
            (pid, clean_date(r.get("dateDepot")), clean_str(r.get("commune")),
             clean_str(r.get("preciserAutre")), r.get("nombreChats"),
             r.get("nombreChatons"), r.get("nombreAutres"), uid),
        )
        n += 1
    return n


def migrate_abandon(cur, reg, wb):
    n = 0
    for r in read_sheet(wb, "Abandon"):
        pid = reg.get_or_create_personne(
            nom=r.get("nom"), prenom=r.get("prenom"), adresse=r.get("adresse"),
            code_postal=r.get("codePostal"), ville=r.get("ville"),
            email=r.get("email"), telephone=r.get("telephone"),
        )
        aid = reg.get_or_create_animal(
            r.get("nomAnimal"), sexe=r.get("sexe"), date_naissance=r.get("dateNaissance"),
            puce=r.get("identification"),
        )
        uid = reg.get_or_create_utilisateur(r.get("saisiPar"))
        cause = "; ".join(x for x in [clean_str(r.get("causeAbandon1")), clean_str(r.get("causeAbandon2"))] if x)
        qualites = "; ".join(x for x in [clean_str(r.get("qualites1")), clean_str(r.get("qualites2"))] if x)
        defauts = "; ".join(x for x in [clean_str(r.get("defauts1")), clean_str(r.get("defauts2"))] if x)
        sante = "; ".join(x for x in [clean_str(r.get("problemeSante1")), clean_str(r.get("problemeSante2"))] if x)
        cur.execute(
            """
            insert into abandons (animal_id, personne_id, cause_abandon, problemes_sante,
                qualites, defauts, vaccins, date_fait, saisi_par)
            values (%s,%s,%s,%s,%s,%s,%s,%s,%s)
            """,
            (aid, pid, cause or None, sante or None, qualites or None, defauts or None,
             clean_str(r.get("vaccins")), clean_date(r.get("dateFait")), uid),
        )
        n += 1
    return n


def migrate_certificat_engagement(cur, reg, wb):
    n = 0
    for r in read_sheet(wb, "Certificat_Engagement"):
        pid = reg.get_or_create_personne(
            nom=r.get("nomAdoptant"), prenom=r.get("prenomAdoptant"),
            adresse=r.get("adresseAdoptant"), code_postal=r.get("codePostalAdoptant"),
            ville=r.get("villeAdoptant"), email=r.get("email"), telephone=r.get("telephone"),
        )
        uid = reg.get_or_create_utilisateur(r.get("saisiPar"))
        cur.execute(
            "insert into certificats_engagement (personne_id, date_fait, saisi_par) "
            "values (%s,%s,%s)",
            (pid, clean_date(r.get("dateCertificat")), uid),
        )
        n += 1
    return n


def migrate_famille_accueil(cur, reg, wb, sheet_name, type_accueil):
    n = 0
    for r in read_sheet(wb, sheet_name):
        pid = reg.get_or_create_personne(
            nom=r.get("nom"), prenom=r.get("prenom"), adresse=r.get("adresse"),
            code_postal=r.get("code_postal"), ville=r.get("ville"),
            email=r.get("email"), telephone=r.get("telephone"),
        )
        nom_animal = r.get("nom_animal") or r.get("nom_chat")
        aid = reg.get_or_create_animal(
            nom_animal, sexe=r.get("sexe"), couleur=r.get("couleur"),
            date_naissance=r.get("date_naissance"), puce=r.get("puce"),
        )
        uid = reg.get_or_create_utilisateur(r.get("saisiPar"))
        cur.execute(
            """
            insert into familles_accueil (type_accueil, personne_id, animal_id,
                date_debut, date_fin, date_fait, saisi_par)
            values (%s,%s,%s,%s,%s,%s,%s)
            """,
            (type_accueil, pid, aid, clean_date(r.get("date_debut")),
             clean_date(r.get("date_fin")), clean_date(r.get("date_fait")), uid),
        )
        n += 1
    return n


def migrate_reservation(cur, reg, wb, sheet_name, type_reservation):
    n = 0
    for r in read_sheet(wb, sheet_name):
        nom_chat = r.get("nomChat") or r.get("animal_autre_texte")
        aid = reg.get_or_create_animal(
            nom_chat, sexe=None, couleur=r.get("couleur"),
            date_naissance=r.get("dateNaissance"), puce=r.get("identification"),
            nom_maman=r.get("nomMaman"), signes_particuliers=r.get("signesParticuliers"),
        )
        pid = reg.get_or_create_personne(
            nom=r.get("nom"), prenom=r.get("prenom"), adresse=r.get("adresseAdoptant"),
            ville=r.get("ville"), code_postal=r.get("codePostal"),
            email=r.get("email"), telephone=r.get("telephone"),
            date_naissance=r.get("dateNaissancePersonne"),
        )
        uid = reg.get_or_create_utilisateur(r.get("saisiPar"))
        montant = sum(clean_amount(r.get(c)) or 0 for c in
                      ("montant_cb", "montant_espece", "montant_virement", "montant_cheque")) or None
        cur.execute(
            """
            insert into reservations (type_reservation, animal_id, personne_id,
                numero_box, date_reservation, montant, montant_virement,
                montant_espece, montant_cb, signes_particuliers, superficie, saisi_par)
            values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            returning id
            """,
            (type_reservation, aid, pid, clean_num_text(r.get("numeroBox")),
             clean_date(r.get("dateReservation")), montant,
             clean_amount(r.get("montant_virement")), clean_amount(r.get("montant_espece")),
             clean_amount(r.get("montant_cb")), clean_str(r.get("signesParticuliers")),
             clean_amount(r.get("superficie")), uid),
        )
        resa_id = cur.fetchone()[0]
        insert_cheques(cur, "reservations_cheques", "reservation_id", resa_id, r)
        if r.get("Photo_Animal_URL") and aid:
            cur.execute(
                "insert into animaux_photos (animal_id, url, est_photo_principale) "
                "values (%s,%s,true) on conflict do nothing",
                (aid, clean_str(r.get("Photo_Animal_URL"))),
            )
        n += 1
    return n


def migrate_adoption(cur, reg, wb):
    n = 0
    for r in read_sheet(wb, "Adoption"):
        aid = reg.get_or_create_animal(
            r.get("nomUsuel"), nom_adoption=r.get("nomAdoption"), sexe=r.get("sexe"),
            couleur=r.get("couleur"), date_naissance=r.get("dateNaissance"),
            puce=r.get("puce"),
        )
        pid = reg.get_or_create_personne(
            nom=r.get("nom"), prenom=r.get("prenom"), adresse=r.get("adresse"),
            code_postal=r.get("codePostal"), ville=r.get("ville"),
            email=r.get("email"), telephone=r.get("portable"),
        )
        uid = reg.get_or_create_utilisateur(r.get("saisiPar"))
        cur.execute(
            """
            insert into adoptions (animal_id, personne_id, date_adoption, nom_attestation,
                tarif_particulier, total_participation, montant_virement, montant_espece,
                montant_cb, date_fait, saisi_par)
            values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            returning id
            """,
            (aid, pid, clean_date(r.get("dateAdoption")), clean_str(r.get("nomAttestation")),
             clean_amount(r.get("tarifParticulier")), clean_amount(r.get("totalParticipation")),
             clean_amount(r.get("montant_virement")), clean_amount(r.get("montant_espece")),
             clean_amount(r.get("montant_cb")), clean_date(r.get("dateFait")), uid),
        )
        adopt_id = cur.fetchone()[0]
        insert_cheques(cur, "adoptions_cheques", "adoption_id", adopt_id, r)
        n += 1
    return n


def migrate_pret_materiel(cur, reg, wb):
    n = 0
    for r in read_sheet(wb, "Pret_Materiel"):
        pid = reg.get_or_create_personne(
            nom=r.get("nomEmprunteur"), prenom=r.get("prenomEmprunteur"),
            adresse=r.get("adresseEmprunteur"), email=r.get("emailEmprunteur"),
            telephone=r.get("telephoneEmprunteur"),
        )
        materiel = clean_str(r.get("materiel")) or clean_str(r.get("materiel_autre_texte"))
        cur.execute(
            "insert into materiel_mouvements (type_mouvement, personne_id, materiel, date_mouvement) "
            "values ('pret',%s,%s,%s)",
            (pid, materiel, clean_date(r.get("datePriseEnCharge"))),
        )
        n += 1
    return n


def migrate_icad(cur, reg, wb):
    n = 0
    for r in read_sheet(wb, "ICAD_Suivi"):
        cle = clean_num_text(r.get("Cle"))
        # La clé ICAD source (n° de puce, avec ou sans espaces) n'est pas
        # toujours réconciliable avec un animal déjà migré depuis les autres
        # onglets (puce parfois absente là-bas) : on tente par puce, sinon on
        # ignore la ligne (rien de plus fiable sans intervention manuelle).
        cur.execute("select id from animaux where puce = %s limit 1", (cle,))
        row = cur.fetchone()
        if not row:
            continue
        cur.execute(
            """
            insert into icad_suivi (animal_id, date_transfert, commentaire)
            values (%s,%s,%s)
            on conflict (animal_id) do nothing
            """,
            (row[0], clean_date(r.get("DateTransfert")), clean_str(r.get("Commentaire"))),
        )
        n += 1
    return n


def migrate_compteurs(cur, wb):
    n = 0
    for r in read_sheet(wb, "Compteurs"):
        m = re.match(r"RecuFiscal(Particulier|Entreprise)_(\d{4})", str(r.get("Compteur") or ""))
        if not m:
            continue
        type_ = "particulier" if m.group(1) == "Particulier" else "entreprise"
        annee = int(m.group(2))
        cur.execute(
            """
            insert into compteurs_recus_fiscaux (type, annee, dernier_numero)
            values (%s,%s,%s)
            on conflict (type, annee) do update set dernier_numero = excluded.dernier_numero
            """,
            (type_, annee, int(r.get("Valeur") or 0)),
        )
        n += 1
    return n


# ============================================================
# Point d'entrée
# ============================================================

def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    conn = psycopg2.connect(DATABASE_URL)
    cur = conn.cursor()
    reg = Registry(cur)

    counts = {}
    counts["adhesions"] = migrate_adhesion(cur, reg, wb)
    counts["dons"] = migrate_dons(cur, reg, wb)
    counts["depots_chat"] = migrate_depot_chat(cur, reg, wb)
    counts["abandons"] = migrate_abandon(cur, reg, wb)
    counts["certificats_engagement"] = migrate_certificat_engagement(cur, reg, wb)
    counts["familles_accueil (provisoire)"] = migrate_famille_accueil(
        cur, reg, wb, "Famille_Accueil_Provisoire", "provisoire")
    counts["familles_accueil (chat_libre)"] = migrate_famille_accueil(
        cur, reg, wb, "Famille_Accueil_Chat_Libre", "chat_libre")
    counts["familles_accueil (adoption)"] = migrate_famille_accueil(
        cur, reg, wb, "Famille_Accueil_Adoption", "adoption")
    counts["reservations (chat)"] = migrate_reservation(
        cur, reg, wb, "Contrat_reservation_chat", "chat")
    counts["reservations (chaton)"] = migrate_reservation(
        cur, reg, wb, "Reservation_Chaton", "chaton")
    counts["adoptions"] = migrate_adoption(cur, reg, wb)
    counts["materiel_mouvements"] = migrate_pret_materiel(cur, reg, wb)
    counts["icad_suivi"] = migrate_icad(cur, reg, wb)
    counts["compteurs_recus_fiscaux"] = migrate_compteurs(cur, wb)

    conn.commit()

    print("=== Migration terminée ===")
    for k, v in counts.items():
        print(f"  {k:35s} {v:4d} lignes")
    print()
    print("=== Personnes / animaux / utilisateurs ===")
    for k, v in reg.stats.items():
        print(f"  {k:25s} {v:4d}")
    print()
    print("/!\\ Comptes utilisateurs créés avec un email temporaire "
          "@arche-import.local — à relier aux vrais comptes une fois "
          "l'authentification Supabase en place, avant mise en production.")

    cur.close()
    conn.close()


if __name__ == "__main__":
    main()
